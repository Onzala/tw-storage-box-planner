# 合併各品牌 CSV（含收納配件 acc_*.csv） → data/boxes_all.csv，並產生檢查用試算表 收納盒資料庫_v<版號>.xlsx
# 用法：python tools/build_db.py 2
import csv, sys, re
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
ver = sys.argv[1] if len(sys.argv) > 1 else '1'
BRAND_FILES = [f for f in ['muji', 'nitori', 'ikea', 'daiso'] if (ROOT / 'data' / f'{f}.csv').exists()]
rows, H = [], None
for f in BRAND_FILES:
    with open(ROOT / 'data' / f'{f}.csv', encoding='utf-8-sig') as fh:
        r = list(csv.DictReader(fh))
    H = H or list(r[0].keys())
    assert list(r[0].keys()) == H, f'{f} 欄位不一致'
    rows += r

# 基本檢查
ids = [r['id'] for r in rows]
assert len(ids) == len(set(ids)), '重複 id'
for r in rows:
    o = [float(r[k]) for k in ('outer_w', 'outer_d', 'outer_h')]
    assert max(o) <= 40, f"{r['id']} 超過 40cm"
    for a, b in zip(('inner_w', 'inner_d', 'inner_h'), o):
        assert not r[a] or float(r[a]) <= b, f"{r['id']} 內寸大於外寸"
    assert r['price_twd'] or r['notes'].startswith('SET:'), f"{r['id']} 沒有價格也不是套組"

with open(ROOT / 'data' / 'boxes_all.csv', 'w', encoding='utf-8-sig', newline='') as fh:
    w = csv.DictWriter(fh, H); w.writeheader(); w.writerows(rows)

# 收納配件：合併 acc_*.csv，類別統一成短名（去掉全形括號裡的說明）
ACC_CATS = ['伸縮棒', '伸縮層架', '鐵網', '網片配件', '抽屜隔板', '小型置物架', '掛勾', '吊掛收納']
accs, AH = [], None
for f in BRAND_FILES:
    p = ROOT / 'data' / f'acc_{f}.csv'
    if not p.exists(): continue
    with open(p, encoding='utf-8-sig') as fh:
        r = list(csv.DictReader(fh))
    if not r: continue
    AH = AH or list(r[0].keys())
    assert list(r[0].keys()) == AH, f'acc_{f} 欄位不一致'
    accs += r
for a in accs:
    a['category'] = a['category'].split('（')[0].split('(')[0].strip()
    assert a['category'] in ACC_CATS, f"{a['id']} 類別不在清單：{a['category']}"
    for k in ('length_min', 'length_max', 'width', 'depth', 'height', 'diameter', 'load_kg', 'price_twd'):
        assert not a[k] or float(a[k]) >= 0, f"{a['id']} {k} 不是數字"
    if a['length_min'] and a['length_max']:
        assert float(a['length_min']) <= float(a['length_max']), f"{a['id']} 伸縮範圍顛倒"
    assert a['price_twd'] or a['notes'].startswith('SET:') or a['confidence'] == 'low', f"{a['id']} 沒有價格"
aids = [a['id'] for a in accs]
assert len(aids) == len(set(aids)), '配件重複 id'
if accs:
    with open(ROOT / 'data' / 'accessories_all.csv', 'w', encoding='utf-8-sig', newline='') as fh:
        w = csv.DictWriter(fh, AH); w.writeheader(); w.writerows(accs)

# 收納技巧：合併 tips_*.csv
SPACES = ['抽屜', '衣櫃', '廚房', '水槽下', '洗手台下', '冰箱', '浴室', '玄關', '桌面', '層架', '壁櫥', '通用']
tips, TH = [], None
for p in sorted((ROOT / 'data').glob('tips_*.csv')):
    if p.name == 'tips_all.csv': continue
    with open(p, encoding='utf-8-sig') as fh:
        r = list(csv.DictReader(fh))
    if not r: continue
    TH = TH or list(r[0].keys())
    assert list(r[0].keys()) == TH, f'{p.name} 欄位不一致'
    tips += r
for t in tips:
    assert t['space'] in SPACES, f"{t['id']} 空間不在清單：{t['space']}"
    assert t['url'].startswith('http'), f"{t['id']} 沒有出處網址"
    assert not t['published_date'] or re.fullmatch(r'\d{4}-\d{2}(-\d{2})?', t['published_date']), f"{t['id']} 日期格式錯誤"
    for a in filter(None, t['uses_accessories'].split('/')):
        assert a in ACC_CATS, f"{t['id']} 配件類別不在清單：{a}"
    assert 20 <= len(t['tip']) <= 300, f"{t['id']} 摘要長度 {len(t['tip'])}"
tids = [t['id'] for t in tips]
assert len(tids) == len(set(tids)), '技巧重複 id'
if tips:
    with open(ROOT / 'data' / 'tips_all.csv', 'w', encoding='utf-8-sig', newline='') as fh:
        w = csv.DictWriter(fh, TH); w.writeheader(); w.writerows(tips)

NUM = {'outer_w', 'outer_d', 'outer_h', 'inner_w', 'inner_d', 'inner_h', 'price_twd'}
ZH = {'id': '編號', 'brand': '品牌', 'series': '系列', 'name': '品名', 'product_code': '商品編號', 'outer_w': '外寬cm', 'outer_d': '外深cm', 'outer_h': '外高cm', 'inner_w': '內寬cm', 'inner_d': '內深cm', 'inner_h': '內高cm', 'material': '材質', 'color': '顏色', 'stackable': '可堆疊', 'has_lid': '附蓋', 'price_twd': '售價NT$', 'price_date': '查價日', 'url': '官網連結', 'source': '來源', 'confidence': '可信度', 'notes': '備註'}
FONT = 'Microsoft JhengHei'
hf = Font(name=FONT, bold=True, color='FFFFFF'); fill = PatternFill('solid', fgColor='4F6D7A'); bf = Font(name=FONT)

wb = Workbook(); ws = wb.active; ws.title = '商品總表'
ws.append([ZH[h] for h in H] + ['套組／多件裝'])
for c in ws[1]: c.font = hf; c.fill = fill; c.alignment = Alignment(horizontal='center')
for r in rows:
    m = re.match(r'SET:([^|]*)\|(\d*)\|(\d+)\|(\d+);', r['notes'])
    setinfo = f'整組 NT${m[2]}，此尺寸 {m[3]} 件／共 {m[4]} 件' if m else ''
    ws.append([(float(r[h]) if r[h] else None) if h in NUM else r[h] for h in H] + [setinfo])
widths = {'name': 38, 'notes': 50, 'url': 30, 'material': 16, 'color': 16, 'series': 18, 'product_code': 16}
for i, h in enumerate(H, 1):
    ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 9 if h in NUM else 11)
ws.column_dimensions[get_column_letter(len(H) + 1)].width = 28
ui = H.index('url')
for row in ws.iter_rows(min_row=2):
    for c in row: c.font = bf
    u = row[ui]
    if u.value: u.hyperlink = u.value; u.font = Font(name=FONT, color='0563C1', underline='single')
ws.freeze_panes = 'E2'; ws.auto_filter.ref = ws.dimensions
n = len(rows) + 1

s = wb.create_sheet('品牌摘要')
s.append(['品牌', '品項數', '有內寸', '單售有價格', '套組拆件／多件裝', '高可信度', '單售平均價NT$'])
for c in s[1]: c.font = hf; c.fill = fill
R = lambda col: f'商品總表!${col}$2:${col}${n}'
brands = [f.upper() for f in BRAND_FILES]
last = len(brands) + 1
for i, b in enumerate(brands, 2):
    s.append([b, f'=COUNTIF({R("B")},A{i})', f'=COUNTIFS({R("B")},A{i},{R("I")},">0")',
              f'=COUNTIFS({R("B")},A{i},{R("P")},">0")', f'=COUNTIFS({R("B")},A{i},{R("V")},"整組*")',
              f'=COUNTIFS({R("B")},A{i},{R("T")},"high")', f'=IFERROR(ROUND(AVERAGEIF({R("B")},A{i},{R("P")}),0),"-")'])
s.append(['合計'] + [f'=SUM({c}2:{c}{last})' for c in 'BCDEF'] + [f'=ROUND(AVERAGE({R("P")}),0)'])
s.append([])
s.append(['資料查詢日 2026-10-26。來源：MUJI shop.muji.tw＋官方收納型錄 25.12；NITORI nitori-net.tw；IKEA ikea.com.tw；大創見各列「來源」欄。'])
s.append(['套組／多件裝的售價欄留空，整組價格寫在「套組／多件裝」欄；計算引擎會依整組件數換算購買組數。'])
for row in s.iter_rows(min_row=2):
    for c in row: c.font = bf
for col, wd in zip('ABCDEFG', [10, 10, 10, 12, 16, 12, 14]): s.column_dimensions[col].width = wd

d = wb.create_sheet('欄位說明')
d.append(['欄位', '說明'])
for c in d[1]: c.font = hf; c.fill = fill
desc = {'outer_w': '外寸寬（正面左右）', 'outer_d': '外寸深（前後）', 'outer_h': '外寸高', 'inner_w': '內寸寬，僅官方有標才填，不推算', 'inner_d': '內寸深，同上', 'inner_h': '內寸高，同上',
        'stackable': 'Y＝官方明確說可堆疊；空白＝官方未說明', 'has_lid': 'Y＝含蓋；另售蓋寫在備註', 'price_twd': '查價日的官網單售價；套組與多件裝留空',
        'confidence': 'high＝官網商品頁直接取得；medium＝型錄或來源間數值衝突（見備註）', 'color': '多色以「/」分隔；單一顏色名含斜線時改用「+」',
        'notes': '套組列開頭為 SET:<套組編號>|<整組價>|<此尺寸件數>|<整組件數>; 供程式讀取'}
for h in H: d.append([ZH[h], desc.get(h, '')])
d.column_dimensions['A'].width = 12; d.column_dimensions['B'].width = 80
for row in d.iter_rows(min_row=2):
    for c in row: c.font = bf

if accs:
    AZH = {'id': '編號', 'brand': '品牌', 'category': '類別', 'name': '品名', 'product_code': '商品編號', 'length_min': '最短cm', 'length_max': '最長cm',
           'width': '寬cm', 'depth': '深cm', 'height': '高cm', 'diameter': '棒徑／線徑cm', 'load_kg': '耐重kg', 'mount': '安裝方式', 'material': '材質',
           'color': '顏色', 'price_twd': '售價NT$', 'price_date': '查價日', 'url': '官網連結', 'source': '來源', 'confidence': '可信度',
           'usage': '官方建議用途', 'notes': '備註'}
    ANUM = {'length_min', 'length_max', 'width', 'depth', 'height', 'diameter', 'load_kg', 'price_twd'}
    wa = wb.create_sheet('收納配件', 1)
    wa.append([AZH.get(h, h) for h in AH])
    for c in wa[1]: c.font = hf; c.fill = fill; c.alignment = Alignment(horizontal='center')
    for a in accs: wa.append([(float(a[h]) if a[h] else None) if h in ANUM else a[h] for h in AH])
    aw = {'name': 38, 'notes': 45, 'usage': 40, 'url': 30, 'category': 12, 'mount': 12, 'material': 14, 'color': 14}
    for i, h in enumerate(AH, 1): wa.column_dimensions[get_column_letter(i)].width = aw.get(h, 9 if h in ANUM else 11)
    ui = AH.index('url')
    for row in wa.iter_rows(min_row=2):
        for c in row: c.font = bf
        if row[ui].value: row[ui].hyperlink = row[ui].value; row[ui].font = Font(name=FONT, color='0563C1', underline='single')
    wa.freeze_panes = 'E2'; wa.auto_filter.ref = wa.dimensions

if tips:
    TZH = {'id': '編號', 'space': '空間', 'topic': '主題', 'title': '技巧', 'tip': '摘要', 'how_to': '步驟', 'uses_products': '提到的盒子系列',
           'uses_accessories': '用到的配件', 'brand_related': '相關品牌', 'source_name': '出處', 'source_type': '來源類型', 'url': '網址',
           'published_date': '發布日', 'checked_date': '查核日', 'lang': '原文語言', 'confidence': '可信度'}
    wt = wb.create_sheet('收納技巧', 2 if accs else 1)
    wt.append([TZH.get(h, h) for h in TH])
    for c in wt[1]: c.font = hf; c.fill = fill; c.alignment = Alignment(horizontal='center')
    for t in tips: wt.append([t[h] for h in TH])
    tw = {'title': 30, 'tip': 60, 'how_to': 50, 'source_name': 36, 'url': 30, 'topic': 16, 'uses_accessories': 16, 'uses_products': 16}
    for i, h in enumerate(TH, 1): wt.column_dimensions[get_column_letter(i)].width = tw.get(h, 11)
    ui = TH.index('url')
    for row in wt.iter_rows(min_row=2):
        for c in row: c.font = bf; c.alignment = Alignment(wrap_text=True, vertical='top')
        if row[ui].value: row[ui].hyperlink = row[ui].value; row[ui].font = Font(name=FONT, color='0563C1', underline='single')
    wt.freeze_panes = 'D2'; wt.auto_filter.ref = wt.dimensions

wb.calculation.fullCalcOnLoad = True
out = ROOT / f'收納盒資料庫_v{ver}.xlsx'
wb.save(out)
from collections import Counter
print('rows', len(rows), Counter(r['brand'] for r in rows), 'sets', sum(r['notes'].startswith('SET:') for r in rows), '->', out.name)
print('accessories', len(accs), Counter(a['brand'] for a in accs), Counter(a['category'] for a in accs))
print('tips', len(tips), Counter(t['source_type'] for t in tips), Counter(t['space'] for t in tips))
